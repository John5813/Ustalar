import logging
import os
import uuid

from pptx import Presentation

from . import config
from .image_client import generate_image
from .layouts import SLIDE_H, SLIDE_W, render_canvas
from .models import Brief

log = logging.getLogger("renderer")


def build_presentation(brief: Brief) -> str:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank_layout = prs.slide_layouts[6]  # to'liq bo'sh layout

    for s in brief.slides:
        slide = prs.slides.add_slide(blank_layout)

        # image elementlari uchun rasm generatsiyasi
        image_paths: dict[int, str] = {}
        for el in s.canvas.elements:
            if el.type == "image" and el.prompt:
                path = generate_image(el.prompt)
                if path:
                    image_paths[id(el)] = path
                else:
                    log.info("Slayd %s: rasm topilmadi, image elementi o'tkazib yuboriladi", s.index)

        render_canvas(slide, s, image_paths)

    os.makedirs(config.WORK_DIR, exist_ok=True)
    out_path = os.path.join(config.WORK_DIR, f"ppt_{uuid.uuid4().hex[:10]}.pptx")
    prs.save(out_path)
    log.info("PPTX saqlandi: %s", out_path)
    return out_path
